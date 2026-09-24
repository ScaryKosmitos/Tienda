"""
Exporta los datos de la tienda a un archivo de Excel (.xlsx) con cuatro hojas:
Ventas (del período elegido), Más vendidos (del mismo período), Inventario y
Fiado (estado actual). Los números se guardan como números, para que en Excel se
puedan sumar, filtrar u ordenar.

El libro se escribe en el modo "solo escritura" de openpyxl: cada fila va directo
al archivo en vez de quedar en memoria. Con años de ventas es la diferencia entre
unos segundos y varios minutos (y más de 1 GB de memoria). A cambio, las filas se
agregan en orden, de arriba abajo, y el ancho de las columnas se fija antes.
"""
from datetime import datetime

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import base_datos as db
from formato import describir_periodo
from recibo import numero_recibo

ESTILO_TITULO = Font(size=14, bold=True)
ESTILO_ENCABEZADO = Font(bold=True, color="FFFFFF")
FONDO_ENCABEZADO = PatternFill("solid", fgColor="1F538D")
ESTILO_TOTAL = Font(bold=True)
ESTILO_ANULADA = Font(color="888888", italic=True)
ESTILO_STOCK_BAJO = Font(color="D32F2F", bold=True)
CENTRADO = Alignment(horizontal="center")
A_LA_DERECHA = Alignment(horizontal="right")

FORMATO_UNIDADES = "#,##0"
FORMATO_FECHA = "yyyy-mm-dd hh:mm"
FORMATO_PORCENTAJE = "0.0%"


def _formato_pesos(valor):
    """Sin decimales si el valor es entero ($1.500), con dos si no ($1.500,50)."""
    return '"$"#,##0' if float(valor).is_integer() else '"$"#,##0.00'


def _celda(hoja, valor, formato=None, fuente=None, alineacion=None, relleno=None):
    """
    Celda lista para agregar a una fila. Un texto que empieza con '=' (ej: un
    producto llamado '=Promo 2x1') se guarda como texto: si no, Excel lo tomaría
    como fórmula y diría que el archivo está dañado.
    """
    celda = WriteOnlyCell(hoja, value=valor)
    if isinstance(valor, str) and valor.startswith("="):
        celda.data_type = "s"
    if formato:
        celda.number_format = formato
    if fuente:
        celda.font = fuente
    if alineacion:
        celda.alignment = alineacion
    if relleno:
        celda.fill = relleno
    return celda


def _preparar_hoja(hoja, titulo, encabezados, anchos):
    """Fija el ancho de las columnas y escribe el título en la fila 1 y los
    encabezados en la fila 3. Los datos empiezan en la fila 4."""
    for columna, ancho in enumerate(anchos, start=1):
        hoja.column_dimensions[get_column_letter(columna)].width = ancho
    # Los encabezados quedan siempre visibles al bajar por la hoja
    hoja.freeze_panes = "A4"
    hoja.append([_celda(hoja, titulo, fuente=ESTILO_TITULO)])
    hoja.append([])
    hoja.append([
        _celda(hoja, texto, fuente=ESTILO_ENCABEZADO, alineacion=CENTRADO, relleno=FONDO_ENCABEZADO)
        for texto in encabezados
    ])


def _terminar_hoja(hoja, cantidad_filas, cantidad_columnas):
    """Agrega los filtros de Excel a los encabezados y deja una fila en blanco antes de los totales."""
    if cantidad_filas:
        hoja.auto_filter.ref = f"A3:{get_column_letter(cantidad_columnas)}{cantidad_filas + 3}"
    hoja.append([])


def _fila_total(hoja, columna_valor, texto, valor):
    """Fila de total en negrita. El texto va alineado a la derecha en la celda anterior
    al valor, así se extiende hacia las celdas vacías de la izquierda sin cortarse."""
    hoja.append([None] * (columna_valor - 2) + [
        _celda(hoja, texto, fuente=ESTILO_TOTAL, alineacion=A_LA_DERECHA),
        _celda(hoja, valor, _formato_pesos(valor), ESTILO_TOTAL),
    ])


def _ancho(textos, minimo=10, maximo=45):
    """Ancho de columna aproximado según el texto más largo."""
    return max(minimo, min(maximo, max((len(str(t)) for t in textos), default=0) + 3))


def _hoja_ventas(hoja, desde, hasta, generado):
    ventas = db.obtener_ventas(desde, hasta)
    encabezados = ("ID Venta", "Recibo", "Fecha y Hora", "Producto", "Cantidad", "Total", "Pago", "Fiado a")
    _preparar_hoja(hoja, f"Ventas: {describir_periodo(desde, hasta)} (generado {generado})", encabezados, (
        10, 10, 18, _ancho({v["nombre_producto"] for v in ventas} | {"Producto"}), 11, 14, 10,
        _ancho({v["cliente"] or "" for v in ventas} | {"Fiado a"}),
    ))

    total_vendido = total_fiado = 0.0
    # obtener_ventas las da de la más reciente a la más antigua; en Excel se
    # leen mejor en orden cronológico
    for venta in reversed(ventas):
        anulada, total = venta["anulada"], venta["total"]
        gris = ESTILO_ANULADA if anulada else None
        hoja.append([
            _celda(hoja, venta["id"], fuente=gris, alineacion=CENTRADO),
            _celda(hoja, numero_recibo(venta["recibo_id"]) if venta["recibo_id"] else "—",
                   fuente=gris, alineacion=CENTRADO),
            _celda(hoja, datetime.fromisoformat(venta["fecha"]), FORMATO_FECHA, gris),
            _celda(hoja, venta["nombre_producto"], fuente=gris),
            _celda(hoja, venta["cantidad"], FORMATO_UNIDADES, gris),
            _celda(hoja, total, _formato_pesos(total), gris),
            _celda(hoja, "Anulada" if anulada else (venta["medio"] or "Fiado"), fuente=gris, alineacion=CENTRADO),
            _celda(hoja, venta["cliente"] or "", fuente=gris),
        ])
        if not anulada:
            total_vendido += total
            if venta["cliente"]:
                total_fiado += total

    _terminar_hoja(hoja, len(ventas), len(encabezados))
    _fila_total(hoja, 6, "Total vendido (sin anuladas)", total_vendido)
    if total_fiado:
        _fila_total(hoja, 6, "De eso, fiado", total_fiado)
    return len(ventas)


def _hoja_mas_vendidos(hoja, desde, hasta, generado):
    ranking = db.obtener_mas_vendidos(desde, hasta)
    encabezados = ("#", "Producto", "Unidades Vendidas", "Total", "% de lo Vendido")
    _preparar_hoja(hoja, f"Más vendidos: {describir_periodo(desde, hasta)} (generado {generado})", encabezados, (
        6, _ancho([r[0] for r in ranking] + ["Producto"]), 19, 14, 17,
    ))

    total_general = sum(total for _, _, total in ranking)
    for puesto, (nombre, unidades, total) in enumerate(ranking, start=1):
        hoja.append([
            _celda(hoja, puesto, alineacion=CENTRADO),
            _celda(hoja, nombre),
            _celda(hoja, unidades, FORMATO_UNIDADES),
            _celda(hoja, total, _formato_pesos(total)),
            _celda(hoja, total / total_general if total_general else 0, FORMATO_PORCENTAJE),
        ])
    _terminar_hoja(hoja, len(ranking), len(encabezados))


def _hoja_inventario(hoja, generado, stock_bajo):
    productos = db.buscar_productos()
    encabezados = ("ID", "Código", "Nombre", "Categoría", "Precio", "Stock", "Valor en Stock")
    _preparar_hoja(hoja, f"Inventario (generado {generado})", encabezados, (
        7, 16, _ancho([p["nombre"] for p in productos] + ["Nombre"]),
        _ancho([p["categoria"] for p in productos] + ["Categoría"]), 13, 10, 17,
    ))

    valor_total = 0.0
    for p in productos:
        precio, stock = p["precio"], p["stock"]
        valor = precio * stock
        valor_total += valor
        hoja.append([
            _celda(hoja, p["id"], alineacion=CENTRADO),
            # El código va como texto, para que Excel no le quite los ceros iniciales
            _celda(hoja, p["codigo_barras"] or "", "@", alineacion=CENTRADO),
            _celda(hoja, p["nombre"]),
            _celda(hoja, p["categoria"]),
            _celda(hoja, precio, _formato_pesos(precio)),
            _celda(hoja, stock, FORMATO_UNIDADES, ESTILO_STOCK_BAJO if stock < stock_bajo else None),
            _celda(hoja, valor, _formato_pesos(valor)),
        ])

    _terminar_hoja(hoja, len(productos), len(encabezados))
    _fila_total(hoja, 7, "Valor total del inventario", valor_total)
    return len(productos)


def _hoja_fiado(hoja, generado):
    clientes = db.obtener_clientes()
    encabezados = ("Cliente", "Debe", "Último Movimiento")
    _preparar_hoja(hoja, f"Fiado (generado {generado})", encabezados, (
        _ancho([c["nombre"] for c in clientes] + ["Cliente"]), 14, 19,
    ))

    total_fiado = 0.0
    for cliente in clientes:
        debe, ultimo = cliente["debe"], cliente["ultimo"]
        if debe > 0:
            total_fiado += debe
        hoja.append([
            _celda(hoja, cliente["nombre"]),
            _celda(hoja, debe, _formato_pesos(debe), ESTILO_STOCK_BAJO if debe > 0 else None),
            _celda(hoja, datetime.fromisoformat(ultimo) if ultimo else "—", FORMATO_FECHA, alineacion=CENTRADO),
        ])

    _terminar_hoja(hoja, len(clientes), len(encabezados))
    _fila_total(hoja, 2, "Total que deben", total_fiado)


def exportar_excel(ruta, desde=None, hasta=None, stock_bajo=5):
    """
    Crea el archivo .xlsx en 'ruta'. 'desde' y 'hasta' son fechas 'AAAA-MM-DD'
    opcionales para las hojas de ventas. Los productos con menos de
    'stock_bajo' unidades se marcan en rojo en el inventario.
    Retorna (cantidad_de_ventas, cantidad_de_productos).
    """
    generado = datetime.now().strftime("%Y-%m-%d %H:%M")
    libro = Workbook(write_only=True)

    cantidad_ventas = _hoja_ventas(libro.create_sheet("Ventas"), desde, hasta, generado)
    _hoja_mas_vendidos(libro.create_sheet("Más vendidos"), desde, hasta, generado)
    cantidad_productos = _hoja_inventario(libro.create_sheet("Inventario"), generado, stock_bajo)
    _hoja_fiado(libro.create_sheet("Fiado"), generado)

    libro.save(ruta)
    return cantidad_ventas, cantidad_productos
